# Neutral fixture documents for the Office mockups (no personal content).
import sys
from docx import Document
from docx.shared import Pt
import openpyxl
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.util import Inches, Pt as PPt
out = sys.argv[1]
d = Document()
d.add_heading('Community Garden — Spring Plan', 0)
d.add_paragraph('Prepared for the neighbourhood association, March meeting.')
d.add_heading('Summary', 1)
d.add_paragraph('This spring we will open twelve new raised beds, repair the east fence, and run four Saturday workshops. '
                'The plan below lists costs, volunteers and dates. Everything fits inside the grant we received in January.')
d.add_heading('Budget', 1)
t = d.add_table(rows=1, cols=3); t.style = 'Light Grid Accent 1'
t.rows[0].cells[0].text, t.rows[0].cells[1].text, t.rows[0].cells[2].text = 'Item', 'Quantity', 'Cost'
for a, b, c in [('Raised bed kits', '12', '$1,440'), ('Soil and compost', '8 m³', '$620'), ('Fence repair', '1', '$380'), ('Workshop supplies', '4', '$240')]:
    r = t.add_row().cells; r[0].text, r[1].text, r[2].text = a, b, c
d.add_heading('Workshops', 1)
for w in ['Seed starting — 5 April', 'Composting basics — 12 April', 'Watering without waste — 19 April', 'Kids planting day — 26 April']:
    d.add_paragraph(w, style='List Bullet')
d.add_paragraph('Questions go to the garden committee. Thank you to every volunteer who signed up.')
d.save(f'{out}/Garden plan.docx')

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Budget'
ws.append(['Category', 'January', 'February', 'March', 'Total'])
rows = [('Seeds', 120, 80, 150), ('Tools', 300, 0, 45), ('Soil', 200, 220, 200), ('Water', 60, 55, 70), ('Workshops', 0, 0, 240)]
for i, (c, a, b, m) in enumerate(rows, start=2):
    ws.append([c, a, b, m, f'=SUM(B{i}:D{i})'])
ws.append(['Total', '=SUM(B2:B6)', '=SUM(C2:C6)', '=SUM(D2:D6)', '=SUM(E2:E6)'])
for cell in ws[1]: cell.font = Font(bold=True); cell.fill = PatternFill('solid', fgColor='DDEBF7')
for cell in ws[7]: cell.font = Font(bold=True)
ws.column_dimensions['A'].width = 16
for col in 'BCDE':
    ws.column_dimensions[col].width = 12
    for r in range(2, 8): ws[f'{col}{r}'].number_format = '$#,##0'
v = wb.create_sheet('Volunteers'); v.append(['Name', 'Saturdays', 'Role'])
for r in [('Sam', 4, 'Beds'), ('Priya', 3, 'Workshops'), ('Jordan', 2, 'Fence'), ('Alex', 4, 'Compost')]: v.append(list(r))
wb.save(f'{out}/Garden budget.xlsx')

p = Presentation(); p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
s = p.slides.add_slide(p.slide_layouts[0]); s.shapes.title.text = 'Spring at the Garden'; s.placeholders[1].text = 'Neighbourhood association · March meeting'
for title, bullets in [('What we will build', ['Twelve raised beds', 'A repaired east fence', 'A shared tool shed shelf']),
                       ('Workshops', ['Seed starting', 'Composting basics', 'Watering without waste', 'Kids planting day']),
                       ('How to help', ['Sign up for one Saturday', 'Bring a friend', 'Share the plan'])]:
    sl = p.slides.add_slide(p.slide_layouts[1]); sl.shapes.title.text = title
    tf = sl.placeholders[1].text_frame; tf.text = bullets[0]
    for b in bullets[1:]: tf.add_paragraph().text = b
p.save(f'{out}/Garden talk.pptx')
