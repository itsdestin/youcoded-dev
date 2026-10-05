# Compare each original sample with the copy the editor saved back.
# The trial typed one marker word ("YCTRIAL"), so the only expected difference
# is that word; anything else lost (formulas, tables, images, styles) is a fidelity problem.
import sys, os, zipfile
from docx import Document
import openpyxl
from pptx import Presentation

S = os.path.dirname(os.path.abspath(__file__)) + '/..'
MARK = 'YCTRIAL'

def docx_stats(p):
    d = Document(p)
    text = '\n'.join(par.text for par in d.paragraphs)
    cells = sum(len(r.cells) for t in d.tables for r in t.rows)
    return {
        'paragraphs': len(d.paragraphs), 'tables': len(d.tables), 'table_cells': cells,
        'images': len(d.inline_shapes), 'sections': len(d.sections),
        'styles_used': len({par.style.name for par in d.paragraphs}),
        'chars': len(text.replace(MARK, '')), 'has_mark': MARK in text or any(MARK in c.text for t in d.tables for r in t.rows for c in r.cells),
        'bold_runs': sum(1 for par in d.paragraphs for r in par.runs if r.bold),
        'hyperlinks': sum(len(par._p.xpath('.//w:hyperlink')) for par in d.paragraphs),
        'footnotes_part': any(n.endswith('footnotes.xml') for n in zipfile.ZipFile(p).namelist()),
    }

def xlsx_stats(p):
    wf = openpyxl.load_workbook(p, data_only=False)
    wv = openpyxl.load_workbook(p, data_only=True)
    s = {'sheets': len(wf.sheetnames), 'formulas': 0, 'values': 0, 'merged': 0, 'defined_names': len(list(wf.defined_names)),
         'cond_fmt': 0, 'data_validations': 0, 'charts': 0, 'images': 0, 'has_mark': False}
    cached = {}
    for ws in wf.worksheets:
        s['merged'] += len(ws.merged_cells.ranges)
        s['cond_fmt'] += len(ws.conditional_formatting)
        s['data_validations'] += len(ws.data_validations.dataValidation)
        s['charts'] += len(ws._charts); s['images'] += len(ws._images)
        for row in ws.iter_rows():
            for c in row:
                if c.value is None: continue
                if isinstance(c.value, str) and c.value.startswith('='): s['formulas'] += 1
                else: s['values'] += 1
                if c.value == MARK: s['has_mark'] = True
    for ws in wv.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, (int, float)): cached[(ws.title, c.coordinate)] = c.value
    s['_cached'] = cached
    return s

def pptx_stats(p):
    pr = Presentation(p)
    shapes = sum(len(sl.shapes) for sl in pr.slides)
    pics = sum(1 for sl in pr.slides for sh in sl.shapes if sh.shape_type == 13)
    text = ''.join(sh.text_frame.text for sl in pr.slides for sh in sl.shapes if sh.has_text_frame)
    notes = sum(1 for sl in pr.slides if sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip())
    return {'slides': len(pr.slides), 'shapes': shapes, 'pictures': pics, 'chars': len(text.replace(MARK, '')),
            'notes': notes, 'layouts': len(pr.slide_layouts), 'has_mark': MARK in text}

for name in sorted(os.listdir(f'{S}/samples')):
    a, b = f'{S}/samples/{name}', (sys.argv[1] if len(sys.argv) > 1 else f'{S}/rig/out') + f'/saved-{name}'
    if not os.path.exists(b):
        print(name, 'NO SAVED COPY'); continue
    fn = {'docx': docx_stats, 'xlsx': xlsx_stats, 'pptx': pptx_stats}[name.rsplit('.', 1)[1]]
    try:
        o, n = fn(a), fn(b)
    except Exception as e:
        print(name, 'ERROR', e); continue
    extra = ''
    if '_cached' in o:
        oc, nc = o.pop('_cached'), n.pop('_cached')
        common = set(oc) & set(nc)
        diff = [k for k in common if abs((oc[k] or 0) - (nc[k] or 0)) > 1e-6 * max(1, abs(oc[k] or 0))]
        extra = f"  numbers: {len(oc)} orig, {len(nc)} saved, {len(common)} shared, {len(diff)} differ" + (f" e.g. {diff[:3]}" if diff else '')
    changed = {k: (o[k], n[k]) for k in o if o[k] != n[k]}
    print(f"{name}: same={[k for k in o if o[k]==n[k]]}\n   changed={changed}{extra}")
